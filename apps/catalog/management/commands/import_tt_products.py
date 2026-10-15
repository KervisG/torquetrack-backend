"""Crea los productos de marca TorqueTrack desde el Excel de números TT
(`TorqueTrack-NUMEROS-TT-CROSS-REFERENCE.xlsx`, hoja `Numeros TT`).

Sin argumentos lee `data/tt_products.json`, la hoja del Excel guardada en el
repo: así llega al servidor con el deploy y no hay que copiar el Excel.
Cuando el Excel cambia, se regenera en la máquina local con
`import_tt_products <archivo.xlsx> --to-json apps/catalog/data/tt_products.json`.

Cada fila es la versión TorqueTrack de un producto que ya existe (`ID
original`): se copia ese producto (fitment, foto, ficha) y cambian el número de
parte (el número TT), la marca (`TorqueTrack`) y los datos que trae el Excel.
Los números OEM, Reman y Aftermarket del Excel quedan como referencias
cruzadas, así la búsqueda por cualquiera de ellos encuentra el producto TT. El
original no se toca: la tienda muestra los dos.

El Excel no trae precio TT: el producto nace con `PLACEHOLDER_PRICE` y el
precio real se pone después desde el panel. Solo crea: un producto TT que ya
existe se salta, para no pisar el precio puesto en el panel. Todo o nada: si
falta un producto original, no se escribe nada."""

import json
import re
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from apps.catalog.models import Product
from apps.catalog.services import ensure_brands, ensure_category

DATA_FILE = Path(__file__).resolve().parents[2] / "data" / "tt_products.json"
HOUSE_BRAND = "TorqueTrack"
PLACEHOLDER_PRICE = 1
SHEET = "Numeros TT"
TT_NUMBER = "Número TT"
ORIGINAL_ID = "ID original"
REQUIRED_COLUMNS = (TT_NUMBER, ORIGINAL_ID, "Título")
# Columna del Excel -> clave de `data`. Una celda vacía conserva el dato del
# original: el Excel no siempre trae el código de motor.
TEXT_COLUMNS = {
    "Marca": "make",
    "Modelo": "model",
    "Motor": "engine",
    "Código motor": "engineCode",
    "Categoría": "category",
}
# Las referencias cruzadas salen del Excel tal cual, también vacías: son las
# de la pieza TT, no las del original.
CROSS_REFERENCE_COLUMNS = {
    "Cross reference OEM": "oemPart",
    "Cross reference Reman": "remanPart",
    "Cross reference Aftermarket": "aftermarketPart",
}
# Datos del original que no aplican a la copia TT: identidad, galería propia,
# el origen del registro y el tipo (`OEM`), que ya no describe la pieza.
DROPPED_KEYS = ("id", "slug", "images", "applicationIds", "sourceRangeId", "type")
_YEAR = re.compile(r"\d{4}")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def tt_title(title: str) -> str:
    """El título del Excel con la marca TorqueTrack en lugar de la original:
    `... Turbo Actuator - GM OE` pasa a `... Turbo Actuator - TorqueTrack`."""
    base = title.rsplit(" - ", 1)[0].strip() if " - " in title else title.strip()
    return f"{base} - {HOUSE_BRAND}"


def tt_product_id(tt_number: str) -> str:
    return tt_number.strip().lower()


def _check_columns(header) -> None:
    missing = [column for column in REQUIRED_COLUMNS if column not in header]
    if missing:
        raise CommandError(f"Missing columns in '{SHEET}': {', '.join(missing)}")


def read_tt_rows(path: Path) -> list[dict]:
    """Las filas de la hoja, del `.xlsx` o de su copia `.json`."""
    if path.suffix.lower() == ".json":
        try:
            rows = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise CommandError(f"Could not read {path}: {exc}") from exc
        if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
            raise CommandError(f"{path} must be a list of rows")
        for row in rows:
            _check_columns(row)
        return rows
    return _read_workbook(path)


def _read_workbook(path: Path) -> list[dict]:
    try:
        import openpyxl
    except ImportError as exc:
        raise CommandError("openpyxl is required: pip install -r requirements/base.txt") from exc
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise CommandError(f"Could not open {path}: {exc}") from exc
    if SHEET not in workbook.sheetnames:
        raise CommandError(f"Sheet '{SHEET}' not found in {path}")
    rows = workbook[SHEET].iter_rows(values_only=True)
    header = [_text(cell) for cell in next(rows, ())]
    _check_columns(header)
    return [
        dict(zip(header, row, strict=False))
        for row in rows
        if any(_text(cell) for cell in row)
    ]


def build_tt_data(row: dict, original: dict) -> dict:
    data = {key: value for key, value in original.items() if key not in DROPPED_KEYS}
    tt_number = _text(row[TT_NUMBER])
    data.update(
        id=tt_product_id(tt_number),
        partNumber=tt_number,
        manufacturer=HOUSE_BRAND,
        title=tt_title(_text(row["Título"]) or _text(original.get("title"))),
        price=PLACEHOLDER_PRICE,
        compareAt=0,
        internalNotes=f"TorqueTrack version of {_text(row[ORIGINAL_ID])}.",
    )
    for column, key in TEXT_COLUMNS.items():
        if _text(row.get(column)):
            data[key] = _text(row[column])
    for column, key in CROSS_REFERENCE_COLUMNS.items():
        if column in row:
            data[key] = _text(row[column])
    years = [int(year) for year in _YEAR.findall(_text(row.get("Años")))]
    if years:
        data["yearFrom"], data["yearTo"] = min(years), max(years)
    return data


class Command(BaseCommand):
    help = "Create the TorqueTrack-brand products from the TT numbers spreadsheet."

    def add_arguments(self, parser):
        parser.add_argument(
            "file",
            type=Path,
            nargs="?",
            default=DATA_FILE,
            help="TT numbers .xlsx or its .json copy (default: data/tt_products.json)",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Validate the file and print the summary without writing.",
        )
        parser.add_argument(
            "--to-json",
            type=Path,
            help="Save the sheet rows as JSON to this path and exit (no database).",
        )

    def handle(self, *args, **options):
        rows = read_tt_rows(options["file"])
        if options["to_json"]:
            content = json.dumps(rows, indent=2, ensure_ascii=False, default=str)
            options["to_json"].write_text(content + "\n", encoding="utf-8")
            self.stdout.write(f"Saved {len(rows)} rows to {options['to_json']}")
            return
        problems = []
        seen = set()
        for index, row in enumerate(rows, start=2):
            tt_number = _text(row[TT_NUMBER])
            if not tt_number or not _text(row[ORIGINAL_ID]):
                problems.append(f"row {index}: {TT_NUMBER} and {ORIGINAL_ID} are required")
            elif tt_product_id(tt_number) in seen:
                problems.append(f"row {index}: {tt_number} is repeated")
            seen.add(tt_product_id(tt_number))
        original_ids = {_text(row[ORIGINAL_ID]) for row in rows}
        originals = {
            product.pk: product
            for product in Product.objects.filter(pk__in=original_ids).prefetch_related(
                "applications"
            )
        }
        problems += [
            f"original product not found: {product_id}"
            for product_id in sorted(original_ids - set(originals))
            if product_id
        ]
        if problems:
            raise CommandError("Nothing was written:\n" + "\n".join(f"  - {p}" for p in problems))

        existing = set(Product.objects.filter(pk__in=seen).values_list("id", flat=True))
        new_rows = [row for row in rows if tt_product_id(_text(row[TT_NUMBER])) not in existing]
        self.stdout.write(
            f"TorqueTrack products to create: {len(new_rows)} "
            f"(skipped {len(rows) - len(new_rows)} existing)"
        )
        if options["dry_run"]:
            self.stdout.write("Dry run: nothing was written.")
            return

        with transaction.atomic():
            now = timezone.now()
            for row in new_rows:
                original = originals[_text(row[ORIGINAL_ID])]
                data = build_tt_data(row, original.data or {})
                product = Product(id=data["id"], active=True, updated_at=now)
                product.data = data
                product.save()
                product.applications.set(original.applications.all())
                ensure_category(data.get("category"))
            ensure_brands(HOUSE_BRAND)

        self.stdout.write(f"TorqueTrack products created: {len(new_rows)}")
        self.stdout.write(
            f"Set their real price in the panel: each one starts at ${PLACEHOLDER_PRICE}."
        )
